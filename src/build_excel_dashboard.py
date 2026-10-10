"""
build_excel_dashboard.py
Builds outputs/HR_Attrition_Dashboard.xlsx from the cleaned data.

Every number on the Dashboard sheet is a live Excel formula (COUNTIF / COUNTIFS / AVERAGEIF)
that points at the Data sheet, so the workbook updates itself if the data changes.

Run from the project folder:
    python src/build_excel_dashboard.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from prepare_data import PROJECT_ROOT, attrition_by, clean, load_raw

OUTPUT_FILE = PROJECT_ROOT / "outputs" / "HR_Attrition_Dashboard.xlsx"

# Columns copied to the Data sheet (raw name -> header shown in Excel)
DATA_COLUMNS = {
    "EmployeeNumber": "Employee ID",
    "Age": "Age",
    "AgeBand": "Age Band",
    "Gender": "Gender",
    "MaritalStatus": "Marital Status",
    "Department": "Department",
    "JobRole": "Job Role",
    "JobLevel": "Job Level",
    "MonthlyIncome": "Monthly Income",
    "IncomeBand": "Income Band",
    "OverTime": "Overtime",
    "BusinessTravel": "Business Travel",
    "YearsAtCompany": "Years at Company",
    "TenureBand": "Tenure Band",
    "JobSatisfactionLabel": "Job Satisfaction",
    "EnvironmentSatisfactionLabel": "Environment Satisfaction",
    "WorkLifeBalanceLabel": "Work-Life Balance",
    "Attrition": "Attrition",
}

# Breakdown tables on the Dashboard. `order` fixes the row order for banded/ranked categories;
# None means "sort by attrition rate, highest first".
BREAKDOWNS = [
    ("Overtime", "OverTime", ["Yes", "No"]),
    ("Job Role", "JobRole", None),
    ("Department", "Department", None),
    ("Age Band", "AgeBand", ["18-25", "26-35", "36-45", "46-60"]),
    ("Tenure Band", "TenureBand", ["0-2 yrs", "3-5 yrs", "6-10 yrs", "10+ yrs"]),
    ("Income Band", "IncomeBand", ["Q1 (Lowest)", "Q2", "Q3", "Q4 (Highest)"]),
    ("Job Satisfaction", "JobSatisfactionLabel", ["Low", "Medium", "High", "Very High"]),
    ("Work-Life Balance", "WorkLifeBalanceLabel", ["Bad", "Good", "Better", "Best"]),
    ("Business Travel", "BusinessTravel", ["Travel Frequently", "Travel Rarely", "Non-Travel"]),
    ("Marital Status", "MaritalStatus", None),
]

# Styling
FONT = "Arial"
NAVY = "1F3864"
LIGHT = "D9E2F3"
title_font = Font(name=FONT, size=18, bold=True, color=NAVY)
subtitle_font = Font(name=FONT, size=10, italic=True, color="595959")
section_font = Font(name=FONT, size=12, bold=True, color=NAVY)
header_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
body_font = Font(name=FONT, size=10)
kpi_label_font = Font(name=FONT, size=9, bold=True, color="595959")
kpi_value_font = Font(name=FONT, size=20, bold=True, color=NAVY)
header_fill = PatternFill("solid", fgColor=NAVY)
kpi_fill = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_data_sheet(wb, df):
    """Write the cleaned employee table and return {raw column name: column letter}."""
    ws = wb.create_sheet("Data")
    letters = {}
    for col_idx, (raw_name, header) in enumerate(DATA_COLUMNS.items(), start=1):
        letter = get_column_letter(col_idx)
        letters[raw_name] = letter
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font, cell.fill, cell.alignment = header_font, header_fill, center
        ws.column_dimensions[letter].width = max(12, len(header) + 4)

    for row_idx, record in enumerate(df[list(DATA_COLUMNS)].itertuples(index=False), start=2):
        for col_idx, value in enumerate(record, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value).font = body_font

    income_col = letters["MonthlyIncome"]
    for row in range(2, len(df) + 2):
        ws[f"{income_col}{row}"].number_format = "#,##0"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DATA_COLUMNS))}{len(df) + 1}"
    return letters


def data_range(letters, raw_name, n_rows):
    col = letters[raw_name]
    return f"Data!${col}$2:${col}${n_rows + 1}"


def write_kpis(ws, letters, n_rows):
    attr = data_range(letters, "Attrition", n_rows)
    income = data_range(letters, "MonthlyIncome", n_rows)
    kpis = [
        ("TOTAL EMPLOYEES", f"=COUNTA({attr})", "#,##0"),
        ("EMPLOYEES WHO LEFT", f'=COUNTIF({attr},"Yes")', "#,##0"),
        ("ATTRITION RATE", "=IFERROR(C5/A5,0)", "0.0%"),
        ("AVG MONTHLY INCOME - STAYED", f'=IFERROR(AVERAGEIF({attr},"No",{income}),0)', "#,##0"),
        ("AVG MONTHLY INCOME - LEFT", f'=IFERROR(AVERAGEIF({attr},"Yes",{income}),0)', "#,##0"),
    ]
    # Each KPI card spans two columns: A:B, C:D, E:F, G:H, I:J
    for i, (label, formula, fmt) in enumerate(kpis):
        first = 1 + i * 2
        last = first + 1
        ws.merge_cells(start_row=4, start_column=first, end_row=4, end_column=last)
        ws.merge_cells(start_row=5, start_column=first, end_row=5, end_column=last)
        label_cell = ws.cell(row=4, column=first, value=label)
        value_cell = ws.cell(row=5, column=first, value=formula)
        label_cell.font, value_cell.font = kpi_label_font, kpi_value_font
        value_cell.number_format = fmt
        for row in (4, 5):
            for col in (first, last):
                c = ws.cell(row=row, column=col)
                c.fill, c.alignment, c.border = kpi_fill, center, box
    ws.row_dimensions[5].height = 36


def write_breakdown(ws, df, letters, n_rows, start_row, title, raw_col, order):
    """One table (Category | Employees | Left | Attrition Rate) plus a bar chart beside it."""
    categories = order or list(attrition_by(df, raw_col).index)
    cat_range = data_range(letters, raw_col, n_rows)
    attr = data_range(letters, "Attrition", n_rows)

    ws.cell(row=start_row, column=1, value=f"Attrition by {title}").font = section_font
    header_row = start_row + 1
    for col, text in enumerate([title, "Employees", "Left", "Attrition Rate"], start=1):
        c = ws.cell(row=header_row, column=col, value=text)
        c.font, c.fill, c.alignment, c.border = header_font, header_fill, center, box

    first_data = header_row + 1
    for i, category in enumerate(categories):
        r = first_data + i
        values = [
            category,
            f"=COUNTIF({cat_range},A{r})",
            f'=COUNTIFS({cat_range},A{r},{attr},"Yes")',
            f"=IFERROR(C{r}/B{r},0)",
        ]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.font, c.border = body_font, box
            if col > 1:
                c.alignment = Alignment(horizontal="center")
            if col in (2, 3):
                c.number_format = "#,##0"
        ws.cell(row=r, column=4).number_format = "0.0%"
    last_data = first_data + len(categories) - 1

    ws.conditional_formatting.add(
        f"D{first_data}:D{last_data}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=0.5, color="C00000"),
    )

    chart = BarChart()
    chart.type = "bar"  # horizontal bars
    chart.style = 10
    chart.title = f"Attrition Rate by {title}"
    chart.legend = None
    chart.add_data(Reference(ws, min_col=4, min_row=header_row, max_row=last_data), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first_data, max_row=last_data))
    chart.y_axis.number_format = "0%"
    chart.y_axis.majorGridlines = None
    chart.x_axis.scaling.orientation = "maxMin"  # keep the same top-to-bottom order as the table
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.series[0].graphicalProperties.solidFill = NAVY
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showLegendKey = False
    chart.height = max(7.0, 1.0 * len(categories) + 2.5)
    chart.width = 15.0
    ws.add_chart(chart, f"F{start_row}")

    # Leave enough rows for whichever is taller: the table or the chart (a row is ~0.53 cm)
    chart_rows = int(chart.height / 0.53) + 3
    return start_row + max(len(categories) + 4, chart_rows)


def write_about_sheet(wb, n_rows):
    ws = wb.create_sheet("About")
    ws.column_dimensions["A"].width = 110
    lines = [
        ("About this workbook", section_font),
        ("Built automatically by src/build_excel_dashboard.py (Python + openpyxl).", body_font),
        ("", body_font),
        ("Sheets", section_font),
        ("Dashboard: KPI cards and attrition breakdowns with charts.", body_font),
        (f"Data: the cleaned employee table ({n_rows:,} rows) that every Dashboard formula reads from.", body_font),
        ("", body_font),
        ("How the numbers work", section_font),
        ("Employees = COUNTIF on the category column of the Data sheet.", body_font),
        ('Left = COUNTIFS on the category column AND Attrition = "Yes".', body_font),
        ("Attrition Rate = Left / Employees.", body_font),
        ("Income bands are quartiles of Monthly Income (each band holds about 25% of employees).", body_font),
        ("", body_font),
        ("Data source", section_font),
        ("IBM HR Analytics Employee Attrition & Performance dataset (Kaggle).", body_font),
        ("It is a fictional dataset created by IBM data scientists for learning and practice.", body_font),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text).font = font


def build(output_file: Path = OUTPUT_FILE) -> Path:
    df = clean(load_raw())
    n_rows = len(df)

    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    letters = write_data_sheet(wb, df)

    ws["A1"] = "HR Employee Attrition Dashboard"
    ws["A1"].font = title_font
    ws["A2"] = "Who is leaving the company, and why? All figures are live formulas linked to the Data sheet."
    ws["A2"].font = subtitle_font
    write_kpis(ws, letters, n_rows)

    next_row = 8
    for title, raw_col, order in BREAKDOWNS:
        next_row = write_breakdown(ws, df, letters, n_rows, next_row, title, raw_col, order)

    ws.column_dimensions["A"].width = 26
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 13
    ws.sheet_view.showGridLines = False

    # Print / PDF export: whole dashboard one page wide, landscape
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    write_about_sheet(wb, n_rows)
    wb.calculation.fullCalcOnLoad = True  # Excel computes every formula when the file opens

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    return output_file


if __name__ == "__main__":
    path = build()
    print(f"Dashboard saved to {path.relative_to(PROJECT_ROOT)}")
